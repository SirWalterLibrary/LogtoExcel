# import modules
import csv
import pandas as pd
from sys import exit
from tkinter import Tk, filedialog
from openpyxl.worksheet import cell_range 
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils.dataframe import dataframe_to_rows

class Log:

    def __init__(self):
        # create empty data list
        self.data = []

        # gets headers of headers.txt
        self.headers = self.getHeaders()

    def getData(self):
        # open dialog to browse file in File Explorer 
        win = Tk()
        win.withdraw()
        win.iconbitmap(r'src/log.ico')
        filepath = filedialog.askopenfilename(initialdir='/', 
                                              title="Select a File", 
                                              filetypes=[("Log files", ".txt .log")])
        
        # check if user imported file 
        if filepath != '':
            # append file contents into data list 
            with open(filepath) as logFile:
                for row in csv.reader(logFile, delimiter=';'):
                    self.data.append(row)
        else:
            # error if no file was imported
            print("No log file was imported")
            exit()

    def parseData(self):
        # create data frame and parse through data
        self.getData()
        self.data = pd.DataFrame(self.data).iloc[:,1:-1].dropna()


    def getHeaders(self, line_number=None):
        # specify header file
        filename = 'src/headers.txt'

        # open header file
        with open(filename) as file:
            # get all headers
            if line_number == None:
                return [x.strip(' ') for x in (" ".join(line.strip() for line in file)).split(';')]
            else:
                for i, line in enumerate(file):
                    # get only the headers from the line
                    if i == line_number:
                        return list(filter(None,(line.strip()).split(';')))

class Category(Log):

    def __init__(self, index=None):
        # calls original __init__()
        super().__init__()
        self.headers = self.getHeaders(index)

    def __call__(self,other):
        # get initial headers and its range
        initial_headers = self.getHeaders(0)
        initial_range = [x + 1 for x, header in enumerate(initial_headers)]
        
        # split data to its respective category (i.e. Dimensions, Contour Verify, Corner)
        self.range = initial_range + [index for index, item in enumerate(other.headers) if item in self.headers]
        self.data = (other.data[self.range]).set_axis((initial_headers + self.headers),axis=1)
    
    def formatTable(self, worksheet):
        # define range
        full_range = cell_range.CellRange(min_col=worksheet.min_column,
                                          min_row=worksheet.min_row,
                                          max_col=worksheet.max_column,
                                          max_row=worksheet.max_row).coord

        # set table format
        mediumStyle =TableStyleInfo(name='TableStyleMedium1',
                                    showRowStripes=True)
        # create a table
        table = Table(ref=full_range,
                    displayName=(worksheet.title).replace(" ", "_"),
                    tableStyleInfo=mediumStyle)

        # add the table to the worksheet
        worksheet.add_table(table)

    def paste2Excel(self,worksheet,title):
        # title the worksheet
        worksheet.title = title

        # paste data frame into Excel worksheet
        for rows in dataframe_to_rows(self.data, index=False, header=True):
            worksheet.append(rows)

        # format table from data
        self.formatTable(worksheet)
